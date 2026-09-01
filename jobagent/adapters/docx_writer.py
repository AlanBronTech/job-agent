"""Rendering documents to Alan's .docx format.

Every measurement here came off `~/Documents/AlanBronResumeMaster2026.docx`,
not from taste. If this file and the master ever disagree, the master wins.

Two decisions worth knowing:

* **A template document is opened rather than a blank one.** The master has no
  `Normal` style and its bullets come from a real numbering definition in
  `numbering.xml`; reconstructing either in python-docx is fragile and would
  drift. `templates/resume_base.docx` is the master with every paragraph and
  table stripped out — styles, numbering and section setup kept, none of
  Alan's text — so the output inherits the format exactly.
* **Formatting is applied per run, not through named styles.** The master does
  it that way, and matching it means a generated document opens beside the
  master with no visible difference.

This module renders. It does not decide what goes in — that is `core/generate`.
"""

from __future__ import annotations

import copy
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph

TEMPLATE_PATH = Path(__file__).resolve().parents[2] / "templates" / "resume_base.docx"

FONT = "Calibri"
LINE_SPACING = 1.1

NAVY = RGBColor(0x1B, 0x3A, 0x6B)      # name, section headings
BLUE = RGBColor(0x22, 0x58, 0xA5)      # tagline
GREY = RGBColor(0x59, 0x59, 0x59)      # contact line
BLACK = RGBColor(0x00, 0x00, 0x00)

# Right-aligned tab the date range on a role heading sits against.
DATE_TAB_CM = 16.933

# The numbering definition the master's bullets use, from its numbering.xml.
BULLET_NUM_ID = 2


class DocxError(Exception):
    """The document could not be rendered."""


# --------------------------------------------------------------------------- #
# What a rendered document is made of
# --------------------------------------------------------------------------- #


@dataclass
class SkillCategory:
    label: str
    skills: str


@dataclass
class RoleBlock:
    title: str
    company: str
    dates: str
    bullets: list[str]


@dataclass
class ResumeContent:
    """Everything a rendered resume needs, already selected and ordered.

    Plain strings by the time they arrive here: selection, ordering and any
    reweighting happened in `core/generate`, and every string traces to
    `profile/`.
    """

    name: str
    tagline: str
    contact: str
    profile_paragraphs: list[str]
    highlights: list[tuple[str, str]]          # (bold label, remainder)
    skills: list[SkillCategory]
    roles: list[RoleBlock]
    earlier_career: list[str]
    education: list[str]


@dataclass
class CoverLetterContent:
    name: str
    contact: str
    date: str
    recipient: list[str] = field(default_factory=list)
    salutation: str = "Dear Hiring Manager,"
    paragraphs: list[str] = field(default_factory=list)
    closing: str = "Regards,"


# --------------------------------------------------------------------------- #
# Rendering
# --------------------------------------------------------------------------- #


def write_resume(content: ResumeContent, path: Path) -> Path:
    """Render a resume and save it. Returns the path written."""
    doc = _open_template()

    _letterhead(doc, content.name, content.tagline, content.contact)

    _section_heading(doc, "PROFILE")
    for index, text in enumerate(content.profile_paragraphs):
        para = _paragraph(doc, text, size=10.5, colour=BLACK)
        para.paragraph_format.space_before = Pt(2.0) if index == 0 else None
        para.paragraph_format.space_after = Pt(3.5) if index == 0 else Pt(2.0)

    if content.highlights:
        _section_heading(doc, "CAREER HIGHLIGHTS")
        for label, rest in content.highlights:
            _bullet_with_label(doc, label, rest)

    if content.skills:
        _section_heading(doc, "CORE SKILLS")
        _skills_table(doc, content.skills)

    _section_heading(doc, "EXPERIENCE")
    for role in content.roles:
        _role_heading(doc, role.title, role.company, role.dates)
        for bullet in role.bullets:
            _bullet(doc, bullet)

    if content.earlier_career:
        _section_heading(doc, "EARLIER CAREER")
        for line in content.earlier_career:
            _bullet(doc, line)

    if content.education:
        _section_heading(doc, "EDUCATION & CERTIFICATIONS")
        for line in content.education:
            _bullet(doc, line)

    return _save(doc, path)


def write_cover_letter(content: CoverLetterContent, path: Path) -> Path:
    """Render a cover letter and save it. Returns the path written."""
    doc = _open_template()

    _letterhead(doc, content.name, None, content.contact)

    _spacer(doc)
    _paragraph(doc, content.date, size=10.5, colour=BLACK)
    for line in content.recipient:
        _paragraph(doc, line, size=10.5, colour=BLACK)

    _spacer(doc)
    _paragraph(doc, content.salutation, size=10.5, colour=BLACK)

    for text in content.paragraphs:
        para = _paragraph(doc, text, size=10.5, colour=BLACK)
        para.paragraph_format.space_before = Pt(6.0)

    _spacer(doc)
    _paragraph(doc, content.closing, size=10.5, colour=BLACK)
    _paragraph(doc, content.name, size=10.5, colour=BLACK)

    return _save(doc, path)


# --------------------------------------------------------------------------- #
# Building blocks
# --------------------------------------------------------------------------- #


def _open_template() -> Document:
    try:
        return Document(str(TEMPLATE_PATH))
    except Exception as exc:
        raise DocxError(
            f"Could not open the document template at {TEMPLATE_PATH}: {exc}"
        ) from exc


def _save(doc: Document, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        doc.save(str(path))
    except OSError as exc:
        raise DocxError(f"Could not write {path}: {exc}") from exc
    return path


def _letterhead(doc: Document, name: str, tagline: str | None, contact: str) -> None:
    heading = _paragraph(doc, name, size=26, bold=True, colour=NAVY)
    heading.paragraph_format.space_after = Pt(1.5)

    if tagline:
        line = _paragraph(doc, tagline, size=10.5, colour=BLUE)
        line.paragraph_format.space_after = Pt(2.0)

    _paragraph(doc, contact, size=9.5, colour=GREY)


def _section_heading(doc: Document, text: str) -> None:
    para = _paragraph(doc, text.upper(), size=11.5, bold=True, colour=NAVY)
    para.paragraph_format.space_before = Pt(11.5)
    para.paragraph_format.space_after = Pt(5.5)


def _role_heading(doc: Document, title: str, company: str, dates: str) -> None:
    """`Title  ·  Company` with the dates right-aligned on a tab stop.

    Three sizes in one line, as the master has it: the title carries the
    weight at 11.5, the company sits at 10.5, and the dates recede to 9.5.
    """
    para = doc.add_paragraph()
    _format_paragraph(para, space_before=Pt(9.0), space_after=Pt(1.0))
    para.paragraph_format.tab_stops.add_tab_stop(
        Cm(DATE_TAB_CM), WD_TAB_ALIGNMENT.RIGHT
    )

    _run(para, title, size=11.5, bold=True)
    if company:
        _run(para, "  ·  ", size=10.5)
        _run(para, company, size=10.5, bold=True)
    if dates:
        _run(para, f"\t{dates}", size=9.5)


def _bullet(doc: Document, text: str) -> Paragraph:
    para = doc.add_paragraph(style="List Paragraph")
    _apply_bullet_numbering(para)
    _format_paragraph(para, space_before=Pt(1.5), space_after=Pt(3.6))
    _run(para, text, size=10.5)
    return para


def _bullet_with_label(doc: Document, label: str, rest: str) -> Paragraph:
    """A CAREER HIGHLIGHTS bullet: bold label, em dash, then the claim."""
    para = doc.add_paragraph(style="List Paragraph")
    _apply_bullet_numbering(para)
    _format_paragraph(para, space_before=Pt(1.5), space_after=Pt(3.6))
    _run(para, f"{label} — ", size=10.5, bold=True)
    _run(para, rest, size=10.5)
    return para


def _skills_table(doc: Document, skills: list[SkillCategory]) -> None:
    """Borderless two-column table, categories filling left-to-right."""
    rows = (len(skills) + 1) // 2
    table = doc.add_table(rows=rows, cols=2)
    table.autofit = False
    _remove_table_borders(table)

    usable = doc.sections[0].page_width - doc.sections[0].left_margin - doc.sections[0].right_margin
    for row in table.rows:
        for cell in row.cells:
            cell.width = int(usable / 2)

    for index, category in enumerate(skills):
        cell = table.rows[index // 2].cells[index % 2]
        # A new cell arrives with one empty paragraph; use it rather than
        # leaving a blank line above the label.
        label = cell.paragraphs[0]
        _format_paragraph(label, space_before=Pt(2.0), space_after=Pt(1.0))
        _run(label, category.label.upper(), size=9.5, bold=True)

        body = cell.add_paragraph()
        _format_paragraph(body, space_after=Pt(4.0))
        _run(body, category.skills, size=9.5)


def _spacer(doc: Document) -> None:
    para = doc.add_paragraph()
    _format_paragraph(para, space_after=Pt(6.0))


def _paragraph(
    doc: Document,
    text: str,
    *,
    size: float,
    bold: bool = False,
    colour: RGBColor | None = None,
) -> Paragraph:
    para = doc.add_paragraph()
    _format_paragraph(para)
    _run(para, text, size=size, bold=bold, colour=colour)
    return para


def _format_paragraph(
    para: Paragraph, *, space_before=None, space_after=None
) -> None:
    fmt = para.paragraph_format
    fmt.line_spacing = LINE_SPACING
    fmt.space_before = space_before
    fmt.space_after = space_after


def _run(
    para: Paragraph,
    text: str,
    *,
    size: float,
    bold: bool = False,
    colour: RGBColor | None = None,
):
    run = para.add_run(text)
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    if colour is not None:
        run.font.color.rgb = colour
    return run


def _apply_bullet_numbering(para: Paragraph) -> None:
    """Point the paragraph at the template's bullet numbering definition.

    python-docx has no API for this. The alternative is the built-in "List
    Bullet" style, which the master does not use and which renders at a
    different indent.
    """
    numbering = para._p.get_or_add_pPr().makeelement(qn("w:numPr"), {})
    level = numbering.makeelement(qn("w:ilvl"), {qn("w:val"): "0"})
    num_id = numbering.makeelement(qn("w:numId"), {qn("w:val"): str(BULLET_NUM_ID)})
    numbering.append(level)
    numbering.append(num_id)
    para._p.get_or_add_pPr().append(numbering)


def _remove_table_borders(table) -> None:
    properties = table._tbl.tblPr
    borders = properties.makeelement(qn("w:tblBorders"), {})
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.makeelement(qn(f"w:{edge}"), {qn("w:val"): "none"})
        borders.append(element)
    properties.append(borders)
