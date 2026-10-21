"""Reading a resume or cover letter .docx back into named sections, for `diff`.

The other direction from `docx_writer`, and looser on purpose: the file being
read may have been edited in Word or rewritten in Claude.ai, so nothing about
its formatting can be assumed. Structure is recovered from text alone:

* A **section heading** is a short paragraph in capitals — PROFILE, CORE
  SKILLS. An edited resume that renamed one gets a section of that name, and
  the comparison reports it as new rather than guessing what it replaced.
* Within EXPERIENCE, a **role heading** is a paragraph that is not a list item
  and holds the `·` between title and company. The dates after the tab are
  dropped from its name, so correcting a date does not orphan the role.
* The CORE SKILLS **table** is read cell by cell in document order. Reading
  `doc.paragraphs` alone skips tables entirely.
* The **letterhead**: the first paragraph is the name, whatever its case, and
  the last before the first heading is the contact line. Both come from the
  profile and are not compared; anything between them is the tagline.

A cover letter is one section, the body between the salutation and the
closing. Comparing it paragraph by paragraph would mis-align every paragraph
after an inserted one.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from jobagent.core.diff import Section

_CLOSING = re.compile(
    r"^(regards|kind regards|best regards|warm regards|yours sincerely|"
    r"yours faithfully|sincerely|thank you|thanks)[,.]?$",
    re.IGNORECASE,
)


class DocxReadError(Exception):
    """The file could not be read as a .docx."""


def read_sections(path: Path, kind: str) -> list[Section]:
    """`kind` is `resume` or `cover_letter`."""
    if Path(path).suffix.lower() != ".docx":
        raise DocxReadError(
            f"{Path(path).name} is not a .docx. Compare the .docx it was exported from."
        )
    try:
        doc = Document(str(path))
    except Exception as exc:
        raise DocxReadError(f"Could not read {path}: {exc}") from exc
    blocks = list(_blocks(doc))
    if kind == "cover_letter":
        return [_letter_body([text for text, _ in blocks])]
    return _resume_sections(blocks)


def _blocks(doc):
    """(text, is_list_item) for every non-empty paragraph, tables included, in order."""
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            para = Paragraph(child, doc)
            if para.text.strip():
                yield para.text, _is_list(para)
        elif child.tag == qn("w:tbl"):
            for row in Table(child, doc).rows:
                seen = set()
                for cell in row.cells:
                    # A merged cell is returned once per grid column it spans.
                    if id(cell._tc) in seen:
                        continue
                    seen.add(id(cell._tc))
                    for para in cell.paragraphs:
                        if para.text.strip():
                            yield para.text, True


def _is_list(para: Paragraph) -> bool:
    ppr = para._p.pPr
    if ppr is not None and ppr.find(qn("w:numPr")) is not None:
        return True
    return bool(para.style is not None and para.style.name.startswith("List"))


def _is_heading(text: str) -> bool:
    stripped = text.strip()
    return (
        len(stripped) <= 40
        and any(c.isalpha() for c in stripped)
        and stripped == stripped.upper()
    )


def _resume_sections(blocks: list[tuple[str, bool]]) -> list[Section]:
    letterhead: list[str] = []
    sections: list[Section] = []
    current: Section | None = None
    in_experience = False
    for index, (text, is_list) in enumerate(blocks):
        # The name opens the letterhead, and some resumes set it in capitals.
        if index == 0:
            letterhead.append(text)
            continue
        if not is_list and _is_heading(text):
            name = " ".join(text.split()).upper()
            in_experience = name == "EXPERIENCE"
            current = None if in_experience else Section(name, [])
            if current is not None:
                sections.append(current)
            continue
        if not sections and current is None and not in_experience:
            letterhead.append(text)
            continue
        if in_experience and not is_list and "·" in text:
            role = " ".join(text.split("\t")[0].split())
            current = Section(f"EXPERIENCE: {role}", [])
            sections.append(current)
            continue
        if current is None:
            current = Section("EXPERIENCE", [])
            sections.append(current)
        current.paragraphs.append(text)
    tagline = letterhead[1:-1]
    return ([Section("TAGLINE", tagline)] if tagline else []) + sections


def _letter_body(lines: list[str]) -> Section:
    start = next(
        (i + 1 for i, line in enumerate(lines) if line.strip().lower().startswith("dear ")),
        0,
    )
    end = next(
        (i for i in range(len(lines) - 1, start - 1, -1) if _CLOSING.match(lines[i].strip())),
        len(lines),
    )
    return Section("LETTER", lines[start:end])
