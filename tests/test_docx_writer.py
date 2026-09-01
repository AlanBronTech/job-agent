"""Formatting tests for the .docx renderer.

Every figure asserted here was measured off
`~/Documents/AlanBronResumeMaster2026.docx`. The master is the source of
truth: if it and these tests disagree, the master wins and these change.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from jobagent.adapters.docx_writer import (
    CoverLetterContent,
    ResumeContent,
    RoleBlock,
    SkillCategory,
    write_cover_letter,
    write_resume,
)


def make_content(**overrides) -> ResumeContent:
    data = {
        "name": "ALAN BRON",
        "tagline": "Engineering Leader  ·  Co-Founder  ·  AI Adoption",
        "contact": "Sydney, NSW   ·   0459 000 000   ·   someone@example.com",
        "profile_paragraphs": ["First paragraph.", "Second paragraph."],
        "highlights": [("AI in production", "Deployed agentic tooling to a team.")],
        "skills": [
            SkillCategory("AI & Automation", "Agent development, prompt engineering"),
            SkillCategory("Leadership", "Teams to 15, hiring end-to-end"),
            SkillCategory("Backend", "Java, Python"),
        ],
        "roles": [
            RoleBlock("Software Development Manager", "Easy Signs", "2025 – 2026",
                      ["Led 15 engineers across three product streams."]),
            RoleBlock("AI Engineer & Consultant", "Bron Consulting", "2026 – Present",
                      ["Built an agentic system."]),
        ],
        "earlier_career": ["MLC / NAB — superannuation platform upgrades."],
        "education": ["PRINCE2 Foundation — APM Group"],
    }
    data.update(overrides)
    return ResumeContent(**data)


@pytest.fixture
def rendered(tmp_path: Path) -> Document:
    return Document(str(write_resume(make_content(), tmp_path / "resume.docx")))


def find(doc: Document, prefix: str):
    return next(p for p in doc.paragraphs if p.text.startswith(prefix))


def signature(para):
    run = para.runs[0]
    fmt = para.paragraph_format
    return (
        run.font.size.pt,
        bool(run.font.bold),
        str(run.font.color.rgb) if run.font.color and run.font.color.rgb else None,
        fmt.space_before.pt if fmt.space_before else None,
        fmt.space_after.pt if fmt.space_after else None,
        fmt.line_spacing,
    )


# --------------------------------------------------------------------------- #
# Typography, against the master's measurements
# --------------------------------------------------------------------------- #


def test_the_name_block(rendered) -> None:
    assert signature(find(rendered, "ALAN BRON")) == (26.0, True, "1B3A6B", None, 1.5, 1.1)


def test_the_tagline(rendered) -> None:
    assert signature(find(rendered, "Engineering Leader")) == (
        10.5, False, "2258A5", None, 2.0, 1.1,
    )


def test_the_contact_line(rendered) -> None:
    assert signature(find(rendered, "Sydney, NSW")) == (
        9.5, False, "595959", None, None, 1.1,
    )


def test_section_headings(rendered) -> None:
    for heading in ("PROFILE", "CAREER HIGHLIGHTS", "EXPERIENCE", "EARLIER CAREER"):
        assert signature(find(rendered, heading)) == (
            11.5, True, "1B3A6B", 11.5, 5.5, 1.1,
        ), heading


def test_a_role_heading_carries_three_sizes_and_a_right_tab(rendered) -> None:
    """`Title  ·  Company` with the dates right-aligned at 16.93cm."""
    para = find(rendered, "Software Development Manager")
    sizes = [(run.text, run.font.size.pt, bool(run.font.bold)) for run in para.runs]

    assert sizes[0] == ("Software Development Manager", 11.5, True)
    assert sizes[2] == ("Easy Signs", 10.5, True)
    assert sizes[3] == ("\t2025 – 2026", 9.5, False)
    tabs = para.paragraph_format.tab_stops
    assert round(tabs[0].position.cm, 2) == 16.93


def test_bullets_use_the_template_numbering(rendered) -> None:
    para = find(rendered, "Led 15 engineers")

    assert "numPr" in para._p.xml
    assert para.style.name == "List Paragraph"
    assert signature(para)[3:] == (1.5, 3.6, 1.1)


def test_a_highlight_bullet_has_a_bold_label_then_an_em_dash(rendered) -> None:
    para = find(rendered, "AI in production")

    assert para.runs[0].text == "AI in production — "
    assert para.runs[0].font.bold is True
    assert para.runs[1].font.bold is False


def test_the_page_setup_comes_from_the_template(rendered) -> None:
    section = rendered.sections[0]

    assert round(section.left_margin.cm, 2) == 1.68
    assert round(section.right_margin.cm, 2) == 1.68
    assert round(section.top_margin.cm, 2) == 1.32
    assert round(section.bottom_margin.cm, 2) == 1.23


# --------------------------------------------------------------------------- #
# Structure
# --------------------------------------------------------------------------- #


def test_sections_appear_in_the_specified_order(rendered) -> None:
    headings = [
        p.text for p in rendered.paragraphs
        if p.text.isupper() and p.text not in {"ALAN BRON"} and len(p.text) > 4
    ]

    assert headings == [
        "PROFILE",
        "CAREER HIGHLIGHTS",
        "CORE SKILLS",
        "EXPERIENCE",
        "EARLIER CAREER",
        "EDUCATION & CERTIFICATIONS",
    ]


def test_skills_render_as_a_borderless_two_column_table(rendered) -> None:
    table = rendered.tables[0]

    assert len(table.columns) == 2
    assert len(table.rows) == 2          # three categories fill two rows
    assert "AI & AUTOMATION" in table.rows[0].cells[0].text
    borders = table._tbl.tblPr.findall(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblBorders"
    )
    assert borders, "the table must declare borders explicitly, all set to none"


def test_an_empty_section_is_omitted_entirely(tmp_path: Path) -> None:
    path = write_resume(
        make_content(highlights=[], earlier_career=[], skills=[]),
        tmp_path / "sparse.docx",
    )
    text = "\n".join(p.text for p in Document(str(path)).paragraphs)

    assert "CAREER HIGHLIGHTS" not in text
    assert "EARLIER CAREER" not in text
    assert "CORE SKILLS" not in text
    assert "EXPERIENCE" in text


# --------------------------------------------------------------------------- #
# Cover letter
# --------------------------------------------------------------------------- #


def test_a_cover_letter_renders_its_parts(tmp_path: Path) -> None:
    path = write_cover_letter(
        CoverLetterContent(
            name="Alan Bron",
            contact="Sydney, NSW   ·   someone@example.com",
            date="1 September 2026",
            recipient=["Toshiba Global Commerce Solutions"],
            paragraphs=["First paragraph.", "Second paragraph."],
        ),
        tmp_path / "letter.docx",
    )
    text = [p.text for p in Document(str(path)).paragraphs if p.text.strip()]

    assert text[0] == "ALAN BRON" or text[0] == "Alan Bron"
    assert "1 September 2026" in text
    assert "Toshiba Global Commerce Solutions" in text
    assert "Dear Hiring Manager," in text
    assert "First paragraph." in text
    assert text[-1] == "Alan Bron"
